"""Verificación del paso 8: cuestionarios de diagnóstico vs. la columna 'Score'.

Uso (desde la raíz del repo):
    python scripts/verificar_cuestionarios.py \
        --c1 "...Cuestionario_1_2020B_Responses.xlsx" \
        --c2 "...Cuestionario_2_2020B_Responses.xlsx"

Opciones:
    --c1 ARCHIVO        Cuestionario 1 (40 preguntas: A1..C10).
    --c2 ARCHIVO        Cuestionario 2 (40 preguntas: A11..C20; la columna
                        extra "(A10)" se ignora en el cálculo, solo se suma +1
                        para comparar contra el "Score" de Google Forms).
    --final-algebra ARCHIVO   Examen final de Álgebra (verificación con BD:
                        cuántos alumnos tienen más de un intento y que se
                        guarda el primer intento).
    --final-periodo PERIODO   Periodo del examen final (default: 2024B).
    --db-url URL        URL de BD para la verificación del examen final
                        (default: la de .env).
    --periodo PERIODO   Periodo explícito si no se deduce del nombre (ej. 2020B).

No escribe ni agrega ningún Excel al repo; solo lee los archivos indicados.
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.upload_cuestionario_service import (  # noqa: E402
    CUESTIONARIO_CODES,
    _correos_fila,
    _identificadores_fila,
    _leer_workbook,
    _obtener_timestamp,
    _respuestas_key_default,
    _valor,
    calcular_aciertos_cuestionario,
    detectar_columnas_cuestionario,
)
from app.services.upload_diagnostico_service import (  # noqa: E402
    extract_answer_key_from_raw,
    extraer_ano_timestamp,
    fila_corresponde_periodo,
    ts_sort_key,
)

_RESPUESTAS_KEY = _respuestas_key_default()
_FILENAME_PERIODO_RE = re.compile(r"((?:19|20)\d{2})\s*([ABab])")


def _periodo_desde_nombre(path: str) -> tuple[str | None, int | None]:
    m = _FILENAME_PERIODO_RE.search(Path(path).name)
    if not m:
        return None, None
    anio = int(m.group(1))
    letra = m.group(2).upper()
    return f"{anio}{letra}", anio


def _parse_score(raw: object | None) -> int | None:
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return int(raw) if raw.is_integer() else None
    s = str(raw).strip()
    if not s:
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else None
    except ValueError:
        return None


def _key_identidad(row: tuple, cols: dict) -> tuple[str, ...]:
    emails = _correos_fila(row, cols)
    cuentas, folios = _identificadores_fila(row, cols)
    from app.services.upload_cuestionario_service import parse_usuario

    _, lug = parse_usuario(_valor(row, cols["usuario"]))
    partes = [
        str(x) for x in [*emails, *cuentas, *folios, f"U{lug}" if lug else ""] if x
    ]
    if not partes:
        return ()
    return tuple(partes)


def _aciertos_totales(row: tuple, cols: dict, cuestionario: int) -> tuple[int, int]:
    """(aciertos de las 40 preguntas, puntaje equivalente al 'Score' c2 con extra)."""
    preguntas = {
        code: idx
        for code, idx in cols["preguntas"].items()
        if code in set(CUESTIONARIO_CODES[cuestionario])
    }
    aciertos, _ = calcular_aciertos_cuestionario(
        row, preguntas, _RESPUESTAS_KEY, CUESTIONARIO_CODES[cuestionario]
    )
    total = sum(aciertos.values())
    equivalente = total
    if cuestionario == 2:
        extra_idx = cols["preguntas"].get("A10")
        if extra_idx is not None:
            if extract_answer_key_from_raw(_valor(row, extra_idx)) == "d":
                equivalente += 1
    return total, equivalente


def _verificar_cuestionario(
    path: str, cuestionario: int, periodo_cli: str | None
) -> int:
    periodo, anio_periodo = _periodo_desde_nombre(path)
    if periodo_cli:
        periodo = periodo_cli
    if anio_periodo is None and periodo:
        anio_periodo = int(periodo[:4]) if periodo[:4].isdigit() else None

    headers, rows = _leer_workbook(Path(path).read_bytes())
    cols = detectar_columnas_cuestionario(headers)

    score_idx = cols["score"]
    if score_idx is None:
        print(f"[C{cuestionario}] AVISO: no se encontró columna 'Score' en {path}")

    total_ok = 0
    total_revisadas = 0
    sin_score = 0
    filas_por_anio: dict[str, int] = {}
    fallas: list[tuple[int, int, int]] = []

    intentos_por_clave: dict[tuple, list[dict]] = {}
    filas_periodo = 0

    for idx, row in enumerate(rows):
        ts = _obtener_timestamp(row, cols)
        anio = extraer_ano_timestamp(ts)
        anio_label = str(anio) if anio is not None else "sin fecha"
        filas_por_anio[anio_label] = filas_por_anio.get(anio_label, 0) + 1

        if anio_periodo is not None and anio == anio_periodo:
            filas_periodo += 1
            clave = _key_identidad(row, cols)
            if clave:
                intentos_por_clave.setdefault(clave, []).append({"ts": ts, "idx": idx})

        aciertos, equivalente = _aciertos_totales(row, cols, cuestionario)
        score = _parse_score(_valor(row, score_idx)) if score_idx is not None else None
        total_revisadas += 1
        if score is None:
            sin_score += 1
            continue
        if equivalente == score:
            total_ok += 1
        else:
            fallas.append((idx, aciertos, score))

    pct = (total_ok / total_revisadas) * 100 if total_revisadas else 0.0

    repetidos = 0
    for clave, intentos in intentos_por_clave.items():
        if len(intentos) < 2:
            continue
        intentos.sort(key=lambda i: (ts_sort_key(i["ts"]), i["idx"]))
        repetidos += len(intentos) - 1

    print(f"\n=== Cuestionario {cuestionario}: {Path(path).name} ===")
    print(f"  total_filas            : {len(rows)}")
    for anio, n in sorted(filas_por_anio.items()):
        print(f"  filas {anio:<10}: {n}")
    print(f"  filas del periodo ({periodo})     : {filas_periodo}")
    print(f"  intentos repetidos     : {repetidos}")
    if score_idx is None:
        print(f"  SIN columna Score: {sin_score} filas no comparadas")
    else:
        print(f"  coincidencia con Score : {total_ok}/{total_revisadas} ({pct:.1f}%)")
        for idx, aciertos, score in fallas[:20]:
            print(f"    fila {idx}: aciertos={aciertos} score={score}")
        if len(fallas) > 20:
            print(f"    ... y {len(fallas) - 20} más")
    return 0


def _verificar_final(path: str, periodo: str, db_url: str | None) -> int:
    import openpyxl

    from sqlalchemy import select
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.core.config import settings
    from app.models.resultado_diagnostico import ResultadoDiagnostico
    from app.models.respuesta_correcta_diagnostico import RespuestaCorrectaDiagnostico
    from app.seeds.data.respuestas_diagnostico import DEFAULT_RESPUESTAS
    from app.services.upload_diagnostico_service import (
        COL_CUENTA,
        COL_EMAIL,
        COL_FOLIO,
        COL_TIMESTAMP,
        clasificar_columnas_cuenta_folio,
        find_alumno,
        generate_answer_codes,
        load_all_alumnos,
        normalize_email,
    )

    engine = create_async_engine(db_url or settings.database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)

    async def _run() -> None:
        async with factory() as db:
            email_map, cuenta_map, folio_map, alumno_details = await load_all_alumnos(
                db
            )

            key_rows = (
                (
                    await db.execute(
                        select(RespuestaCorrectaDiagnostico).where(
                            RespuestaCorrectaDiagnostico.materia == "algebra",
                            RespuestaCorrectaDiagnostico.periodo == periodo,
                        )
                    )
                )
                .scalars()
                .all()
            )
            if key_rows:
                respuestas_key = {r.codigo: r.respuesta_correcta for r in key_rows}
            else:
                respuestas_key = dict(DEFAULT_RESPUESTAS["algebra"])

            answer_codes = generate_answer_codes("FA", 20)
            correct_key = [
                respuestas_key.get(code, "").lower() for code in answer_codes
            ]

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            wb.close()

            intentos: dict[int, list[dict]] = {}
            for idx, row in enumerate(rows):
                raw_ts = row[COL_TIMESTAMP] if len(row) > COL_TIMESTAMP else None
                if not fila_corresponde_periodo(raw_ts, periodo):
                    continue
                email = normalize_email(
                    row[COL_EMAIL] if len(row) > COL_EMAIL else None
                )
                cuenta, folio = clasificar_columnas_cuenta_folio(
                    row[COL_CUENTA] if len(row) > COL_CUENTA else None,
                    row[COL_FOLIO] if len(row) > COL_FOLIO else None,
                )
                alumno_id = find_alumno(
                    email, cuenta, folio, email_map, cuenta_map, folio_map
                )
                if alumno_id is None:
                    continue
                answers = [
                    extract_answer_key_from_raw(
                        row[7 + qi] if len(row) > 7 + qi else None
                    )
                    for qi in range(20)
                ]
                intentos.setdefault(alumno_id, []).append(
                    {"ts": raw_ts, "idx": idx, "answers": answers}
                )

            multi = {aid: lst for aid, lst in intentos.items() if len(lst) > 1}

            print(
                f"\n=== Examen final Álgebra: {Path(path).name} (periodo {periodo}) ==="
            )
            print(f"  alumnos con más de un intento: {len(multi)}")

            import json as _json

            mismatches: list[tuple[int, str]] = []
            checked = 0
            no_registro = 0
            for aid, lst in sorted(multi.items()):
                oldest = min(lst, key=lambda a: (ts_sort_key(a["ts"]), a["idx"]))
                stored = (
                    (
                        await db.execute(
                            select(ResultadoDiagnostico).where(
                                ResultadoDiagnostico.alumno_id == aid,
                                ResultadoDiagnostico.periodo == periodo,
                            )
                        )
                    )
                    .scalars()
                    .first()
                )
                if stored is None:
                    no_registro += 1
                    mismatches.append((aid, "sin registro guardado"))
                    continue
                expected = []
                for qi, code in enumerate(answer_codes):
                    is_correct = (
                        oldest["answers"][qi] is not None
                        and oldest["answers"][qi] == correct_key[qi]
                    )
                    expected.append(
                        {
                            "codigo": code,
                            "respuesta": oldest["answers"][qi] or "",
                            "correcta": is_correct,
                        }
                    )
                try:
                    guardado = _json.loads(stored.respuestas_algebra or "[]")
                except ValueError:
                    guardado = None
                checked += 1
                if guardado != expected:
                    mismatches.append(
                        (aid, "el registro guardado NO es el primer intento")
                    )

            print(
                f"  verificados (con registro): {checked}, sin registro: {no_registro}"
            )
            if mismatches:
                print(f"  PROBLEMAS ({len(mismatches)}):")
                for aid, motivo in mismatches[:20]:
                    print(f"    alumno {aid}: {motivo}")
            else:
                print(
                    "  OK: el registro guardado es el primer intento (fecha más antigua)."
                )

    asyncio.run(_run())
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--c1", metavar="ARCHIVO", help="Archivo del Cuestionario 1")
    parser.add_argument("--c2", metavar="ARCHIVO", help="Archivo del Cuestionario 2")
    parser.add_argument("--periodo", default=None, help="Periodo ej. 2020B")
    parser.add_argument(
        "--final-algebra",
        metavar="ARCHIVO",
        help="Examen final de Álgebra (requiere BD)",
    )
    parser.add_argument(
        "--final-periodo", default="2024B", help="Periodo del examen final"
    )
    parser.add_argument("--db-url", default=None, help="URL de BD (default: .env)")
    args = parser.parse_args()

    if not (args.c1 or args.c2 or args.final_algebra):
        parser.error("Se requiere al menos uno de: --c1, --c2, --final-algebra")

    rc = 0
    if args.c1:
        rc |= _verificar_cuestionario(args.c1, 1, args.periodo)
    if args.c2:
        rc |= _verificar_cuestionario(args.c2, 2, args.periodo)
    if args.final_algebra:
        rc |= _verificar_final(args.final_algebra, args.final_periodo, args.db_url)
    return rc


if __name__ == "__main__":
    sys.exit(main())
