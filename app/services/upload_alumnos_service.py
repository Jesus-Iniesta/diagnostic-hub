from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.alumno import Alumno
from app.models.ingenieria import Ingenieria
from app.models.role import Role
from app.models.user import AuthMethod, User
from app.services.normalizacion import normalizar_cuenta, normalizar_folio


# ── Columnas del Excel (índices 0-based) ──────────────────────────
COL_NOMBRE = 19
COL_APELLIDO_PATERNO = 17
COL_APELLIDO_MATERNO = 18
COL_CORREO_PERSONAL = 8
COL_CORREO_INSTITUCIONAL = 26
COL_NUMERO_CUENTA = 13
COL_NUMERO_FOLIO = 12
COL_INGENIERIA = 6
COL_PERIODO = 7
COL_PROMEDIO = 9
COL_INDICE_UAEM = 10
COL_LUGAR_ADMISION = 11
COL_TIENE_INTERNET = 15
COL_TIENE_COMPUTADORA = 16
COL_ES_FORANEO = 23
COL_CONVIVENCIA = 24
COL_VULNERABILIDAD = 27
COL_ESCUELA = 29


@dataclass
class FilaResultado:
    fila: int
    nombre_completo: str
    numero_cuenta: str | None
    estado: str  # "exitoso" | "duplicado" | "error"
    motivo: str = ""
    campos_con_error: list[dict] = field(default_factory=list)
    datos_originales: dict = field(default_factory=dict)


@dataclass
class ResultadoCarga:
    total_filas: int = 0
    exitosos: int = 0
    duplicados: int = 0
    errores: int = 0
    detalle: list[FilaResultado] = field(default_factory=list)


def _clean_str(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_title(value: str) -> str:
    return value.title() if value else ""


def _parse_bool(value: object | None) -> bool:
    return _clean_str(value).lower() in ("sí", "si", "yes", "true", "1")


def _parse_float(value: object | None) -> float | None:
    s = _clean_str(value)
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _parse_int(value: object | None) -> int | None:
    s = _clean_str(value)
    if not s:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def _extract_clave_ingenieria(raw: str) -> str | None:
    match = re.search(r"\((\w+)\)", raw)
    if match:
        return match.group(1).upper()
    cleaned = raw.strip().upper()
    if cleaned:
        return cleaned
    return None


_clean_numero_cuenta = normalizar_cuenta
_clean_numero_folio = normalizar_folio


def _clean_email(value: object | None) -> str | None:
    s = _clean_str(value).lower()
    if not s or "@" not in s:
        return None
    return s


def _build_datos_originales(
    row: tuple,
    *,
    nombre: str,
    ap_paterno: str,
    ap_materno: str,
    correo: str | None,
    correo_inst: str | None,
    num_cuenta: str | None,
    num_folio: str | None,
    ingenieria_raw: str,
    periodo: str,
    promedio: float | None,
    indice: float | None,
    lugar: int | None,
    internet: bool,
    computadora: bool,
    foraneo: bool,
    convivencia: str | None,
    vulnerabilidad: bool,
    escuela: str | None,
) -> dict:
    return {
        "nombre": row[COL_NOMBRE] if row[COL_NOMBRE] else "",
        "apellido_paterno": row[COL_APELLIDO_PATERNO] if row[COL_APELLIDO_PATERNO] else "",
        "apellido_materno": row[COL_APELLIDO_MATERNO] if row[COL_APELLIDO_MATERNO] else "",
        "correo_personal": row[COL_CORREO_PERSONAL] if row[COL_CORREO_PERSONAL] else "",
        "correo_institucional": row[COL_CORREO_INSTITUCIONAL] if row[COL_CORREO_INSTITUCIONAL] else "",
        "numero_cuenta": row[COL_NUMERO_CUENTA] if row[COL_NUMERO_CUENTA] else "",
        "numero_folio": row[COL_NUMERO_FOLIO] if row[COL_NUMERO_FOLIO] else "",
        "ingenieria": ingenieria_raw,
        "periodo": periodo,
        "promedio_bachillerato": row[COL_PROMEDIO] if row[COL_PROMEDIO] else "",
        "indice_uaem": row[COL_INDICE_UAEM] if row[COL_INDICE_UAEM] else "",
        "lugar_admision": row[COL_LUGAR_ADMISION] if row[COL_LUGAR_ADMISION] else "",
        "tiene_internet": "Sí" if internet else "No",
        "tiene_computadora": "Sí" if computadora else "No",
        "es_foraneo": "Sí" if foraneo else "No",
        "convivencia": convivencia or "",
        "vulnerabilidad_economica": "Sí" if vulnerabilidad else "No",
        "escuela": escuela or "",
    }


async def _load_dup_sets(db: AsyncSession) -> tuple[set[str], set[str], set[str]]:
    all_cuentas = await db.execute(
        select(Alumno.numero_cuenta).where(Alumno.numero_cuenta.isnot(None))
    )
    existing_cuentas = {r[0] for r in all_cuentas}

    all_folios = await db.execute(
        select(Alumno.numero_folio).where(Alumno.numero_folio.isnot(None))
    )
    existing_folios = {r[0] for r in all_folios}

    all_correos = await db.execute(
        select(User.correo_personal).where(User.correo_personal.isnot(None))
    )
    existing_correos = {r[0] for r in all_correos}

    return existing_cuentas, existing_folios, existing_correos


async def _get_role_and_ingenierias(db: AsyncSession) -> tuple[Role, dict[str, int]]:
    role = await db.scalar(select(Role).where(Role.name == "alumno"))
    if not role:
        raise ValueError("Rol 'alumno' no encontrado en la DB")

    ingenierias_rows = await db.execute(select(Ingenieria))
    ingenierias_map = {i.clave: i.id for i in ingenierias_rows.scalars()}

    return role, ingenierias_map


def _validate_row(
    *,
    nombre: str,
    ap_paterno: str,
    ap_materno: str,
    correo: str | None,
    num_cuenta: str | None,
    num_folio: str | None,
    ingenieria_raw: str,
    ingenierias_map: dict[str, int],
    periodo: str,
    nombre_completo: str,
    fila_num: int,
) -> list[dict]:
    campos: list[dict] = []

    if not nombre:
        campos.append({"campo": "nombre", "motivo": "Nombre es obligatorio"})
    if not ap_paterno:
        campos.append({"campo": "apellido_paterno", "motivo": "Apellido paterno es obligatorio"})
    if not ap_materno:
        campos.append({"campo": "apellido_materno", "motivo": "Apellido materno es obligatorio"})
    if not correo:
        campos.append({"campo": "correo_personal", "motivo": "Correo personal inválido o vacío"})

    if not num_cuenta and not num_folio:
        campos.append({"campo": "numero_cuenta", "motivo": "Se requiere número de cuenta o número de folio"})
        campos.append({"campo": "numero_folio", "motivo": "Se requiere número de cuenta o número de folio"})
    elif num_folio and len(num_folio) != 9:
        campos.append({"campo": "numero_folio", "motivo": f"Folio inválido (se requieren 9 dígitos, tiene {len(num_folio)})"})

    clave = _extract_clave_ingenieria(ingenieria_raw)
    if not clave or clave not in ingenierias_map:
        campos.append({"campo": "ingenieria", "motivo": f"Ingeniería no reconocida: '{ingenieria_raw}'"})

    if not periodo or not re.match(r"^\d{4}[AB]$", periodo):
        campos.append({"campo": "periodo", "motivo": f"Periodo de ingreso inválido: '{periodo}'"})

    return campos


def _create_user_and_alumno(
    *,
    role: Role,
    ingenierias_map: dict[str, int],
    nombre: str,
    ap_paterno: str,
    ap_materno: str,
    correo: str | None,
    correo_inst: str | None,
    num_cuenta: str | None,
    num_folio: str | None,
    ingenieria_raw: str,
    periodo: str,
    promedio: float | None,
    indice: float | None,
    lugar: int | None,
    internet: bool,
    computadora: bool,
    foraneo: bool,
    convivencia: str | None,
    vulnerabilidad: bool,
    escuela: str | None,
) -> tuple[User, Alumno]:
    clave = _extract_clave_ingenieria(ingenieria_raw)
    assert clave is not None
    ingenieria_id = ingenierias_map[clave]

    user = User(
        nombre=nombre,
        apellido_paterno=ap_paterno,
        apellido_materno=ap_materno,
        correo_personal=correo,
        correo_institucional=correo_inst,
        auth_method=AuthMethod.NUMERO_CUENTA,
        activo=True,
        role_id=role.id,
    )

    alumno = Alumno(
        usuario_id=0,
        ingenieria_id=ingenieria_id,
        numero_cuenta=num_cuenta,
        numero_folio=num_folio,
        periodo_ingreso=periodo,
        promedio_bachillerato=promedio,
        indice_uaem=indice,
        lugar_admision=lugar,
        escuela_procedencia=escuela,
        tiene_internet=internet,
        tiene_computadora=computadora,
        es_foraneo=foraneo,
        convivencia=convivencia,
        vulnerabilidad_economica=vulnerabilidad,
    )

    return user, alumno


async def procesar_excel(db: AsyncSession, file_bytes: bytes) -> ResultadoCarga:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    result = ResultadoCarga(total_filas=len(rows))

    try:
        role, ingenierias_map = await _get_role_and_ingenierias(db)
    except ValueError as e:
        result.errores = result.total_filas
        result.detalle = [
            FilaResultado(fila=i + 2, nombre_completo="", numero_cuenta=None,
                          estado="error", motivo=str(e))
            for i in range(len(rows))
        ]
        return result

    existing_cuentas, existing_folios, existing_correos = await _load_dup_sets(db)
    seen_cuentas: set[str] = set()
    seen_folios: set[str] = set()
    seen_correos: set[str] = set()

    for idx, row in enumerate(rows):
        fila_num = idx + 2

        nombre = _to_title(_clean_str(row[COL_NOMBRE]))
        ap_paterno = _to_title(_clean_str(row[COL_APELLIDO_PATERNO]))
        ap_materno = _to_title(_clean_str(row[COL_APELLIDO_MATERNO]))
        correo = _clean_email(row[COL_CORREO_PERSONAL])
        correo_inst = _clean_email(row[COL_CORREO_INSTITUCIONAL])
        num_cuenta = _clean_numero_cuenta(row[COL_NUMERO_CUENTA])
        num_folio = _clean_numero_folio(row[COL_NUMERO_FOLIO])
        ingenieria_raw = _clean_str(row[COL_INGENIERIA])
        periodo = _clean_str(row[COL_PERIODO])
        promedio = _parse_float(row[COL_PROMEDIO])
        indice = _parse_float(row[COL_INDICE_UAEM])
        lugar = _parse_int(row[COL_LUGAR_ADMISION])
        internet = _parse_bool(row[COL_TIENE_INTERNET])
        computadora = _parse_bool(row[COL_TIENE_COMPUTADORA])
        foraneo = _parse_bool(row[COL_ES_FORANEO])
        convivencia = _clean_str(row[COL_CONVIVENCIA]) or None
        vulnerabilidad = _parse_bool(row[COL_VULNERABILIDAD])
        escuela = _clean_str(row[COL_ESCUELA]) or None

        nombre_completo = f"{nombre} {ap_paterno} {ap_materno}".strip()

        campos_con_error = _validate_row(
            nombre=nombre, ap_paterno=ap_paterno, ap_materno=ap_materno,
            correo=correo, num_cuenta=num_cuenta, num_folio=num_folio,
            ingenieria_raw=ingenieria_raw, ingenierias_map=ingenierias_map,
            periodo=periodo, nombre_completo=nombre_completo, fila_num=fila_num,
        )

        if campos_con_error:
            result.errores += 1
            result.detalle.append(FilaResultado(
                fila=fila_num, nombre_completo=nombre_completo,
                numero_cuenta=num_cuenta, estado="error",
                motivo="; ".join(c["motivo"] for c in campos_con_error),
                campos_con_error=campos_con_error,
                datos_originales=_build_datos_originales(
                    row, nombre=nombre, ap_paterno=ap_paterno, ap_materno=ap_materno,
                    correo=correo, correo_inst=correo_inst, num_cuenta=num_cuenta,
                    num_folio=num_folio, ingenieria_raw=ingenieria_raw, periodo=periodo,
                    promedio=promedio, indice=indice, lugar=lugar, internet=internet,
                    computadora=computadora, foraneo=foraneo, convivencia=convivencia,
                    vulnerabilidad=vulnerabilidad, escuela=escuela,
                ),
            ))
            continue

        is_dup = False
        dup_reason = ""

        if num_cuenta and (num_cuenta in existing_cuentas or num_cuenta in seen_cuentas):
            is_dup = True
            dup_reason = f"Duplicate numero_cuenta '{num_cuenta}'"
        if num_folio and (num_folio in existing_folios or num_folio in seen_folios):
            is_dup = True
            dup_reason = dup_reason or f"Duplicate numero_folio '{num_folio}'"
        if correo in existing_correos or correo in seen_correos:
            is_dup = True
            dup_reason = dup_reason or f"Duplicate correo '{correo}'"

        if is_dup:
            result.duplicados += 1
            result.detalle.append(FilaResultado(
                fila=fila_num, nombre_completo=nombre_completo,
                numero_cuenta=num_cuenta, estado="duplicado",
                motivo=dup_reason,
            ))
            continue

        if num_cuenta:
            seen_cuentas.add(num_cuenta)
        if num_folio:
            seen_folios.add(num_folio)
        if correo:
            seen_correos.add(correo)

        user, alumno = _create_user_and_alumno(
            role=role, ingenierias_map=ingenierias_map,
            nombre=nombre, ap_paterno=ap_paterno, ap_materno=ap_materno,
            correo=correo, correo_inst=correo_inst, num_cuenta=num_cuenta,
            num_folio=num_folio, ingenieria_raw=ingenieria_raw, periodo=periodo,
            promedio=promedio, indice=indice, lugar=lugar, internet=internet,
            computadora=computadora, foraneo=foraneo, convivencia=convivencia,
            vulnerabilidad=vulnerabilidad, escuela=escuela,
        )
        db.add(user)
        await db.flush()
        alumno.usuario_id = user.id
        db.add(alumno)
        result.exitosos += 1
        result.detalle.append(FilaResultado(
            fila=fila_num, nombre_completo=nombre_completo,
            numero_cuenta=num_cuenta, estado="exitoso",
        ))

    await db.commit()
    return result


async def procesar_correcciones(
    db: AsyncSession, filas: list[dict]
) -> ResultadoCarga:
    result = ResultadoCarga(total_filas=len(filas))

    try:
        role, ingenierias_map = await _get_role_and_ingenierias(db)
    except ValueError as e:
        result.errores = result.total_filas
        result.detalle = [
            FilaResultado(fila=f.get("fila", 0), nombre_completo="", numero_cuenta=None,
                          estado="error", motivo=str(e))
            for f in filas
        ]
        return result

    existing_cuentas, existing_folios, existing_correos = await _load_dup_sets(db)
    seen_cuentas: set[str] = set()
    seen_folios: set[str] = set()
    seen_correos: set[str] = set()

    for item in filas:
        fila_num = item.get("fila", 0)
        d = item.get("datos", {})

        nombre = _to_title(_clean_str(d.get("nombre", "")))
        ap_paterno = _to_title(_clean_str(d.get("apellido_paterno", "")))
        ap_materno = _to_title(_clean_str(d.get("apellido_materno", "")))
        correo = _clean_email(d.get("correo_personal", ""))
        correo_inst = _clean_email(d.get("correo_institucional", ""))
        num_cuenta = _clean_numero_cuenta(d.get("numero_cuenta", ""))
        num_folio = _clean_numero_folio(d.get("numero_folio", ""))
        ingenieria_raw = _clean_str(d.get("ingenieria", ""))
        periodo = _clean_str(d.get("periodo", ""))
        promedio = _parse_float(d.get("promedio_bachillerato", ""))
        indice = _parse_float(d.get("indice_uaem", ""))
        lugar = _parse_int(d.get("lugar_admision", ""))
        internet = _parse_bool(d.get("tiene_internet", ""))
        computadora = _parse_bool(d.get("tiene_computadora", ""))
        foraneo = _parse_bool(d.get("es_foraneo", ""))
        convivencia = _clean_str(d.get("convivencia", "")) or None
        vulnerabilidad = _parse_bool(d.get("vulnerabilidad_economica", ""))
        escuela = _clean_str(d.get("escuela", "")) or None

        nombre_completo = f"{nombre} {ap_paterno} {ap_materno}".strip()

        campos_con_error = _validate_row(
            nombre=nombre, ap_paterno=ap_paterno, ap_materno=ap_materno,
            correo=correo, num_cuenta=num_cuenta, num_folio=num_folio,
            ingenieria_raw=ingenieria_raw, ingenierias_map=ingenierias_map,
            periodo=periodo, nombre_completo=nombre_completo, fila_num=fila_num,
        )

        if campos_con_error:
            result.errores += 1
            result.detalle.append(FilaResultado(
                fila=fila_num, nombre_completo=nombre_completo,
                numero_cuenta=num_cuenta, estado="error",
                motivo="; ".join(c["motivo"] for c in campos_con_error),
                campos_con_error=campos_con_error,
            ))
            continue

        is_dup = False
        dup_reason = ""

        if num_cuenta and (num_cuenta in existing_cuentas or num_cuenta in seen_cuentas):
            is_dup = True
            dup_reason = f"Duplicate numero_cuenta '{num_cuenta}'"
        if num_folio and (num_folio in existing_folios or num_folio in seen_folios):
            is_dup = True
            dup_reason = dup_reason or f"Duplicate numero_folio '{num_folio}'"
        if correo in existing_correos or correo in seen_correos:
            is_dup = True
            dup_reason = dup_reason or f"Duplicate correo '{correo}'"

        if is_dup:
            result.duplicados += 1
            result.detalle.append(FilaResultado(
                fila=fila_num, nombre_completo=nombre_completo,
                numero_cuenta=num_cuenta, estado="duplicado",
                motivo=dup_reason,
            ))
            continue

        if num_cuenta:
            seen_cuentas.add(num_cuenta)
        if num_folio:
            seen_folios.add(num_folio)
        if correo:
            seen_correos.add(correo)

        user, alumno = _create_user_and_alumno(
            role=role, ingenierias_map=ingenierias_map,
            nombre=nombre, ap_paterno=ap_paterno, ap_materno=ap_materno,
            correo=correo, correo_inst=correo_inst, num_cuenta=num_cuenta,
            num_folio=num_folio, ingenieria_raw=ingenieria_raw, periodo=periodo,
            promedio=promedio, indice=indice, lugar=lugar, internet=internet,
            computadora=computadora, foraneo=foraneo, convivencia=convivencia,
            vulnerabilidad=vulnerabilidad, escuela=escuela,
        )
        db.add(user)
        await db.flush()
        alumno.usuario_id = user.id
        db.add(alumno)
        result.exitosos += 1
        result.detalle.append(FilaResultado(
            fila=fila_num, nombre_completo=nombre_completo,
            numero_cuenta=num_cuenta, estado="exitoso",
        ))

    await db.commit()
    return result
