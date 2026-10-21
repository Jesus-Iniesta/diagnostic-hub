"""Lista de alumnos del administrador: consulta, edición de datos y el Excel
"Alumnos por revisar"."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.alumno import Alumno
from app.models.fila_no_registrada import FilaNoRegistrada
from app.models.identificador_alumno import IdentificadorAlumno
from app.models.ingenieria import Ingenieria
from app.models.resultado_cuestionario_diagnostico import (
    ResultadoCuestionarioDiagnostico,
)
from app.models.resultado_diagnostico import ResultadoDiagnostico
from app.models.resultado_webassign import ResultadoWebAssign
from app.models.user import User
from app.services.carga_periodo_service import ETIQUETAS, MATERIAS_FINAL
from app.services.normalizacion import normalizar_cuenta, normalizar_folio, normalizar_nombre


class DatosAlumnoError(ValueError):
    pass


def _nombre(user: User) -> str:
    return " ".join(
        p for p in (user.nombre, user.apellido_paterno, user.apellido_materno) if p
    ).strip()


def _clave_periodo(periodo: str | None) -> str:
    from app.services.upload_diagnostico_service import normalizar_periodo

    p = str(periodo or "")
    return normalizar_periodo(p) or "".join(p.upper().split())


async def examenes_por_alumno(
    db: AsyncSession, alumno_ids: list[int], periodo: str | None = None
) -> dict[int, list[str]]:
    """Exámenes con resultado de cada alumno ("Cuestionario 1", "Final Álgebra",
    "WebAssign ICO"...). Sin periodo, se agrega el periodo a cada uno."""
    examenes: dict[int, list[str]] = {aid: [] for aid in alumno_ids}
    if not alumno_ids:
        return examenes

    def agregar(aid: int, etiqueta: str, p: str) -> None:
        examenes[aid].append(etiqueta if periodo is not None else f"{etiqueta} ({p})")

    def filtro(modelo):
        conds = [modelo.alumno_id.in_(alumno_ids)]
        if periodo is not None:
            conds.append(modelo.periodo == periodo)
        return select(modelo).where(*conds).order_by(modelo.periodo)

    for r in (await db.execute(filtro(ResultadoCuestionarioDiagnostico))).scalars():
        for n in (1, 2):
            if any(
                getattr(r, f"aciertos_c{n}_{m}") is not None
                for m in ("algebra", "trigonometria", "geometria", "calculo")
            ):
                agregar(r.alumno_id, f"Cuestionario {n}", r.periodo)
    for r in (await db.execute(filtro(ResultadoDiagnostico))).scalars():
        for m, nombre in MATERIAS_FINAL.items():
            if getattr(r, f"puntaje_{m}") is not None:
                agregar(r.alumno_id, f"Final {nombre}", r.periodo)
    for r in (await db.execute(filtro(ResultadoWebAssign))).scalars():
        agregar(r.alumno_id, f"WebAssign {r.carrera}", r.periodo)
    return examenes


def _fila_alumno(alumno: Alumno, examenes: list[str]) -> dict:
    user = alumno.usuario
    return {
        "id": alumno.id,
        "nombre": _nombre(user),
        "correo_personal": user.correo_personal,
        "correo_institucional": user.correo_institucional,
        "numero_cuenta": alumno.numero_cuenta,
        "numero_folio": alumno.numero_folio,
        "ingenieria_id": alumno.ingenieria_id,
        "ingenieria": alumno.ingenieria.clave if alumno.ingenieria else None,
        "periodo_ingreso": alumno.periodo_ingreso,
        "es_provisional": bool(alumno.es_provisional),
        "examenes": examenes,
    }


def _con_relaciones(query):
    return query.options(selectinload(Alumno.usuario), selectinload(Alumno.ingenieria))


async def listar_alumnos(
    db: AsyncSession,
    *,
    busqueda: str | None,
    periodo: str | None,
    limit: int,
    offset: int,
) -> dict:
    conds = []
    if periodo and periodo.strip():
        clave = _clave_periodo(periodo)
        conds.append(
            func.upper(func.replace(Alumno.periodo_ingreso, " ", "")) == clave
        )
    if busqueda and busqueda.strip():
        term = f"%{busqueda.strip()}%"
        nombre_completo = func.concat_ws(
            " ", User.nombre, User.apellido_paterno, User.apellido_materno
        )
        conds.append(
            or_(
                nombre_completo.ilike(term),
                User.correo_personal.ilike(term),
                User.correo_institucional.ilike(term),
                Alumno.numero_cuenta.ilike(term),
                Alumno.numero_folio.ilike(term),
            )
        )

    base = select(Alumno).join(User, Alumno.usuario_id == User.id).where(*conds)
    total = (
        await db.execute(select(func.count()).select_from(base.subquery()))
    ).scalar_one()
    alumnos = (
        await db.execute(
            _con_relaciones(base)
            # Los registrados sin nombre (provisionales) van al final.
            .order_by(
                User.nombre == "",
                User.nombre,
                User.apellido_paterno,
                User.apellido_materno,
                Alumno.id,
            )
            .offset(offset)
            .limit(limit)
        )
    ).scalars().all()

    examenes = await examenes_por_alumno(db, [a.id for a in alumnos])
    return {
        "items": [_fila_alumno(a, examenes[a.id]) for a in alumnos],
        "total": total,
        "limit": limit,
        "offset": offset,
    }


async def obtener_alumno(db: AsyncSession, alumno_id: int) -> dict | None:
    alumno = await db.scalar(_con_relaciones(select(Alumno).where(Alumno.id == alumno_id)))
    if alumno is None:
        return None
    examenes = await examenes_por_alumno(db, [alumno_id])
    return _fila_alumno(alumno, examenes[alumno_id])


async def _dueno(db: AsyncSession, tipo: str, valor: str, alumno_id: int) -> Alumno | None:
    """Otro alumno que ya tiene esa cuenta/folio (registrada o aprendida)."""
    campo = Alumno.numero_cuenta if tipo == "cuenta" else Alumno.numero_folio
    otro = await db.scalar(
        _con_relaciones(select(Alumno).where(campo == valor, Alumno.id != alumno_id))
    )
    if otro is not None:
        return otro
    otro_id = await db.scalar(
        select(IdentificadorAlumno.alumno_id).where(
            IdentificadorAlumno.tipo == tipo,
            IdentificadorAlumno.valor == valor,
            IdentificadorAlumno.alumno_id != alumno_id,
        )
    )
    if otro_id is None:
        return None
    return await db.scalar(_con_relaciones(select(Alumno).where(Alumno.id == otro_id)))


def _quien(alumno: Alumno) -> str:
    nombre = _nombre(alumno.usuario) or "un alumno sin nombre"
    correo = alumno.usuario.correo_personal or alumno.usuario.correo_institucional
    return f"{nombre} ({correo})" if correo else nombre


async def actualizar_datos(
    db: AsyncSession,
    alumno_id: int,
    *,
    numero_cuenta: str | None,
    numero_folio: str | None,
    ingenieria_id: int | None,
) -> dict:
    """Cambia cuenta, folio e ingeniería. Vacío = quitar el dato.

    Lanza DatosAlumnoError si el formato no es válido o si la cuenta/folio ya
    es de otro alumno (dice de quién) y no guarda nada.
    """
    alumno = await db.scalar(_con_relaciones(select(Alumno).where(Alumno.id == alumno_id)))
    if alumno is None:
        raise LookupError("Alumno no encontrado")

    cuenta = (numero_cuenta or "").strip() or None
    folio = (numero_folio or "").strip() or None
    if cuenta is not None:
        cuenta = normalizar_cuenta(cuenta)
        if cuenta is None:
            raise DatosAlumnoError("El número de cuenta debe tener 7 dígitos.")
    if folio is not None:
        folio = normalizar_folio(folio)
        if folio is None:
            raise DatosAlumnoError("El folio debe tener 9 dígitos.")
    if ingenieria_id is not None and await db.get(Ingenieria, ingenieria_id) is None:
        raise DatosAlumnoError("La ingeniería no existe.")

    for tipo, valor, etiqueta in (("cuenta", cuenta, "La cuenta"), ("folio", folio, "El folio")):
        if valor is None:
            continue
        otro = await _dueno(db, tipo, valor, alumno_id)
        if otro is not None:
            raise DatosAlumnoError(
                f"{etiqueta} {valor} ya la tiene {_quien(otro)}. "
                "Si son la misma persona, únelos."
            )

    alumno.numero_cuenta = cuenta
    alumno.numero_folio = folio
    alumno.ingenieria_id = ingenieria_id
    await db.commit()
    return await obtener_alumno(db, alumno_id)


# ── Excel "Alumnos por revisar" ──────────────────────────────────────

ENCABEZADOS = ["Nombre", "Correo", "Cuenta", "Folio", "Ingeniería", "Exámenes que tiene"]
ANCHOS = [40, 42, 12, 14, 12, 60]


async def _alumnos_del_periodo(db: AsyncSession, periodo: str) -> list[Alumno]:
    """Alumnos que ingresaron en el periodo o tienen algún resultado en él."""
    con_resultado = set()
    for modelo in (ResultadoDiagnostico, ResultadoCuestionarioDiagnostico, ResultadoWebAssign):
        con_resultado |= set(
            (await db.execute(select(modelo.alumno_id).where(modelo.periodo == periodo))).scalars()
        )
    clave = _clave_periodo(periodo)
    alumnos = (await db.execute(_con_relaciones(select(Alumno)))).scalars().all()
    return [
        a for a in alumnos
        if a.id in con_resultado or _clave_periodo(a.periodo_ingreso) == clave
    ]


def _fila_excel(alumno: Alumno, examenes: list[str]) -> list:
    user = alumno.usuario
    return [
        _nombre(user),
        user.correo_personal or user.correo_institucional or "",
        alumno.numero_cuenta or "",
        alumno.numero_folio or "",
        alumno.ingenieria.clave if alumno.ingenieria else "",
        ", ".join(examenes),
    ]


async def casos_por_revisar(db: AsyncSession, periodo: str) -> dict[str, list[list]]:
    """Filas de cada hoja del Excel, ya ordenadas por nombre."""
    periodo = periodo.strip()
    alumnos = await _alumnos_del_periodo(db, periodo)
    examenes = await examenes_por_alumno(db, [a.id for a in alumnos], periodo)

    por_nombre: dict[str, list[Alumno]] = {}
    for a in alumnos:
        clave = normalizar_nombre(_nombre(a.usuario))
        if clave:
            por_nombre.setdefault(clave, []).append(a)
    duplicados = [a for grupo in por_nombre.values() if len(grupo) > 1 for a in grupo]

    no_registradas = (
        await db.execute(select(FilaNoRegistrada).where(FilaNoRegistrada.periodo == periodo))
    ).scalars().all()

    def ordenar(filas: list[list]) -> list[list]:
        return sorted(filas, key=lambda f: (normalizar_nombre(f[0]) or "~", f[0], f[1]))

    return {
        "Posibles duplicados": ordenar([_fila_excel(a, examenes[a.id]) for a in duplicados]),
        "Sin cuenta": ordenar(
            [_fila_excel(a, examenes[a.id]) for a in alumnos if not a.numero_cuenta]
        ),
        "Sin ingeniería": ordenar(
            [_fila_excel(a, examenes[a.id]) for a in alumnos if a.ingenieria_id is None]
        ),
        "No se pudieron registrar": ordenar(
            [
                [f.nombre or "", "", "", "", f.ingenieria or "", ETIQUETAS.get(f.examen, f.examen)]
                for f in no_registradas
            ]
        ),
    }


def generar_excel_por_revisar(casos: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    negritas = Font(bold=True)
    for titulo, filas in casos.items():
        ws = wb.create_sheet(titulo)
        ws.append(ENCABEZADOS)
        for celda in ws[1]:
            celda.font = negritas
        for fila in filas:
            ws.append(fila)
        if not filas:
            ws.append(["Sin alumnos en este caso."])
        for i, ancho in enumerate(ANCHOS):
            # Lo bastante ancho para leer el correo (y el nombre) completo.
            largo = max((len(str(f[i])) for f in filas if i < len(f)), default=0)
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = min(
                max(ancho, largo + 2), 100
            )
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
