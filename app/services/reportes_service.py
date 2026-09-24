import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.alumno import Alumno
from app.models.ingenieria import Ingenieria
from app.models.resultado_diagnostico import ResultadoDiagnostico
from app.models.resultado_webassign import ResultadoWebAssign
from app.models.user import User

HEADER_FILL = PatternFill(start_color="1C1954", end_color="1C1954", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
SUBHEADER_FILL = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, size=10)
DATA_FONT = Font(size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


def _fmt_score(value: float | None) -> float:
    return value if value is not None else 0.0


def _fmt_wa(value: float | None, tiene_wa: bool) -> float | str:
    if not tiene_wa:
        return 0
    if value is None or value == 0:
        return "No trabajó"
    return value


def _style_header(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_data(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        if col > 1:
            cell.alignment = Alignment(horizontal="center")


async def _load_data(
    db: AsyncSession, periodo: str, licenciatura: str | None = None
) -> list[dict]:
    stmt = (
        select(Alumno, User, Ingenieria, ResultadoDiagnostico, ResultadoWebAssign)
        .join(User, Alumno.usuario_id == User.id)
        .outerjoin(Ingenieria, Alumno.ingenieria_id == Ingenieria.id)
        .outerjoin(
            ResultadoDiagnostico,
            (ResultadoDiagnostico.alumno_id == Alumno.id)
            & (ResultadoDiagnostico.periodo == periodo),
        )
        .outerjoin(
            ResultadoWebAssign,
            (ResultadoWebAssign.alumno_id == Alumno.id)
            & (ResultadoWebAssign.periodo == periodo),
        )
        .options(selectinload(Alumno.ingenieria))
    )
    if licenciatura:
        stmt = stmt.where(Ingenieria.clave == licenciatura)
    result = await db.execute(stmt)
    rows = result.all()

    data = []
    seen = set()
    for alumno, user, ingenieria, diag, wa in rows:
        if alumno.id in seen:
            continue
        seen.add(alumno.id)
        data.append({
            "nombre": f"{user.apellido_paterno} {user.apellido_materno} {user.nombre}".strip(),
            "numero_cuenta": alumno.numero_cuenta or "",
            "ingenieria": ingenieria.clave if ingenieria else "",
            "puntaje_algebra": diag.puntaje_algebra if diag else None,
            "puntaje_trigonometria": diag.puntaje_trigonometria if diag else None,
            "puntaje_geometria": diag.puntaje_geometria if diag else None,
            "puntaje_calculo": diag.puntaje_calculo if diag else None,
            "promedio_diagnostico": diag.promedio_diagnostico if diag else None,
            "algebra_trabajo": wa.algebra_trabajo if wa else None,
            "algebra_examen": wa.algebra_examen if wa else None,
            "trigonometria_trabajo": wa.trigonometria_trabajo if wa else None,
            "trigonometria_examen": wa.trigonometria_examen if wa else None,
            "geometria_trabajo": wa.geometria_trabajo if wa else None,
            "geometria_examen": wa.geometria_examen if wa else None,
            "carrera_wa": wa.carrera if wa else "",
            "tiene_wa": wa is not None,
        })

    data.sort(key=lambda r: (r["ingenieria"], r["nombre"]))
    return data


def _write_resumen_ws(ws, data: list[dict]):
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2")
    ws.merge_cells("D1:H1")
    ws.merge_cells("I1:N1")

    ws.cell(1, 1, "Nombre")
    ws.cell(1, 2, "No. de cuenta")
    ws.cell(1, 3, "Lic.")
    ws.cell(1, 4, "Diagnóstico")
    ws.cell(1, 9, "WebAssign")

    sub_headers = [
        "Álgebra", "Trig.", "G. Analítica", "Cálc. Dif", "Promedio",
        "Ál. Trabajo", "Ál. Examen", "Tr. Trabajo", "Tr. Examen",
        "G.A. Trabajo", "G.A. Examen",
    ]
    for i, h in enumerate(sub_headers, 4):
        ws.cell(2, i, h)

    for col in range(1, 15):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        cell = ws.cell(2, col)
        cell.fill = SUBHEADER_FILL
        cell.font = SUBHEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_idx, d in enumerate(data, 3):
        ws.cell(row_idx, 1, d["nombre"])
        ws.cell(row_idx, 2, d["numero_cuenta"])
        ws.cell(row_idx, 3, d["ingenieria"])
        ws.cell(row_idx, 4, _fmt_score(d["puntaje_algebra"]))
        ws.cell(row_idx, 5, _fmt_score(d["puntaje_trigonometria"]))
        ws.cell(row_idx, 6, _fmt_score(d["puntaje_geometria"]))
        ws.cell(row_idx, 7, _fmt_score(d["puntaje_calculo"]))
        ws.cell(row_idx, 8, _fmt_score(d["promedio_diagnostico"]))
        ws.cell(row_idx, 9, _fmt_wa(d["algebra_trabajo"], d["tiene_wa"]))
        ws.cell(row_idx, 10, _fmt_wa(d["algebra_examen"], d["tiene_wa"]))
        ws.cell(row_idx, 11, _fmt_wa(d["trigonometria_trabajo"], d["tiene_wa"]))
        ws.cell(row_idx, 12, _fmt_wa(d["trigonometria_examen"], d["tiene_wa"]))
        ws.cell(row_idx, 13, _fmt_wa(d["geometria_trabajo"], d["tiene_wa"]))
        ws.cell(row_idx, 14, _fmt_wa(d["geometria_examen"], d["tiene_wa"]))
        _style_data(ws, row_idx, 14)

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 8
    for col in range(4, 15):
        ws.column_dimensions[get_column_letter(col)].width = 13


def _write_resumen_final(wb: openpyxl.Workbook, data: list[dict]):
    ws = wb.create_sheet("Resumen final")
    _write_resumen_ws(ws, data)


def _write_resumen_por_licenciatura(
    wb: openpyxl.Workbook,
    data: list[dict],
    licenciaturas: list[tuple[str, str]],
):
    for clave, _ in licenciaturas:
        rows = [d for d in data if d["ingenieria"] == clave]
        if not rows:
            continue
        ws = wb.create_sheet(f"Resumen {clave}")
        _write_resumen_ws(ws, rows)


def _write_diagnostico(wb: openpyxl.Workbook, data: list[dict]):
    ws = wb.create_sheet("Diagnóstico")
    headers = [
        "Nombre", "No. de cuenta", "Licenciatura",
        "Álgebra", "Trigonometría", "Geometría Analítica",
        "Cálculo Diferencial", "Promedio",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for d in data:
        ws.append([
            d["nombre"],
            d["numero_cuenta"],
            d["ingenieria"],
            _fmt_score(d["puntaje_algebra"]),
            _fmt_score(d["puntaje_trigonometria"]),
            _fmt_score(d["puntaje_geometria"]),
            _fmt_score(d["puntaje_calculo"]),
            _fmt_score(d["promedio_diagnostico"]),
        ])
        _style_data(ws, ws.max_row, len(headers))

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    for col in range(4, 9):
        ws.column_dimensions[get_column_letter(col)].width = 16


def _write_webassign(wb: openpyxl.Workbook, data: list[dict]):
    ws = wb.create_sheet("WebAssign")
    headers = [
        "Nombre", "No. de cuenta", "Licenciatura", "Carrera",
        "Ál. Trabajo", "Ál. Examen",
        "Tr. Trabajo", "Tr. Examen",
        "G.A. Trabajo", "G.A. Examen",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    wa_data = [d for d in data if d["algebra_trabajo"] is not None or d["trigonometria_trabajo"] is not None]

    for d in wa_data:
        ws.append([
            d["nombre"],
            d["numero_cuenta"],
            d["ingenieria"],
            d["carrera_wa"],
            _fmt_wa(d["algebra_trabajo"], d["tiene_wa"]),
            _fmt_wa(d["algebra_examen"], d["tiene_wa"]),
            _fmt_wa(d["trigonometria_trabajo"], d["tiene_wa"]),
            _fmt_wa(d["trigonometria_examen"], d["tiene_wa"]),
            _fmt_wa(d["geometria_trabajo"], d["tiene_wa"]),
            _fmt_wa(d["geometria_examen"], d["tiene_wa"]),
        ])
        _style_data(ws, ws.max_row, len(headers))

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    for col in range(5, 11):
        ws.column_dimensions[get_column_letter(col)].width = 14


async def generar_excel_reporte(
    db: AsyncSession, periodo: str, licenciatura: str | None = None
) -> bytes:
    data = await _load_data(db, periodo, licenciatura)

    licenciaturas = sorted({(d["ingenieria"], d["ingenieria"]) for d in data})

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_resumen_final(wb, data)
    _write_resumen_por_licenciatura(wb, data, licenciaturas)
    _write_diagnostico(wb, data)
    _write_webassign(wb, data)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


async def get_periodos_disponibles(db: AsyncSession) -> list[str]:
    stmt_d = select(ResultadoDiagnostico.periodo).distinct()
    stmt_w = select(ResultadoWebAssign.periodo).distinct()
    periods_d = (await db.execute(stmt_d)).scalars().all()
    periods_w = (await db.execute(stmt_w)).scalars().all()
    return sorted(set(periods_d + periods_w), reverse=True)


async def get_stats_reporte(db: AsyncSession, periodo: str) -> dict:
    from sqlalchemy import func

    diag_count = (await db.execute(
        select(func.count()).select_from(ResultadoDiagnostico).where(
            ResultadoDiagnostico.periodo == periodo
        )
    )).scalar_one()

    wa_count = (await db.execute(
        select(func.count()).select_from(ResultadoWebAssign).where(
            ResultadoWebAssign.periodo == periodo
        )
    )).scalar_one()

    wa_by_carrera = (await db.execute(
        select(ResultadoWebAssign.carrera, func.count())
        .where(ResultadoWebAssign.periodo == periodo)
        .group_by(ResultadoWebAssign.carrera)
    )).all()

    diag_by_ing = (await db.execute(
        select(
            Alumno.ingenieria_id,
            func.count(func.distinct(ResultadoDiagnostico.alumno_id)),
        )
        .join(ResultadoDiagnostico, ResultadoDiagnostico.alumno_id == Alumno.id)
        .where(ResultadoDiagnostico.periodo == periodo)
        .group_by(Alumno.ingenieria_id)
    )).all()

    wa_by_ing = (await db.execute(
        select(
            Alumno.ingenieria_id,
            func.count(func.distinct(ResultadoWebAssign.alumno_id)),
        )
        .join(ResultadoWebAssign, ResultadoWebAssign.alumno_id == Alumno.id)
        .where(ResultadoWebAssign.periodo == periodo)
        .group_by(Alumno.ingenieria_id)
    )).all()

    diag_map = dict(diag_by_ing)
    wa_map = dict(wa_by_ing)

    ingenierias = (await db.execute(
        select(Ingenieria).order_by(Ingenieria.nombre)
    )).scalars().all()

    por_licenciatura = [
        {
            "clave": ing.clave,
            "nombre": ing.nombre,
            "diagnosticos": diag_map.get(ing.id, 0),
            "webassign": wa_map.get(ing.id, 0),
        }
        for ing in ingenierias
    ]

    return {
        "periodo": periodo,
        "diagnosticos": diag_count,
        "webassign": wa_count,
        "webassign_por_carrera": {c: n for c, n in wa_by_carrera},
        "por_licenciatura": por_licenciatura,
    }
