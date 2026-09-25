from __future__ import annotations

import io
import json
import re

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.alumno import Alumno
from app.models.resultado_cuestionario_diagnostico import (
    ResultadoCuestionarioDiagnostico,
)
from app.seeds.data.respuestas_cuestionario_diagnostico import (
    DEFAULT_RESPUESTAS_CUESTIONARIO,
)
from app.services.normalizacion import (
    clasificar_identificador,
    normalizar_cuenta,
)
from app.services.upload_diagnostico_service import (
    extract_answer_key_from_raw,
    fila_corresponde_periodo,
    load_all_alumnos,
    normalize_name,
    obtener_rango_periodo,
    rango_fechas_dict,
    ts_sort_key,
)
from app.services.upload_webassign_service import clean_email

MATERIA_FOR_CODE = {
    "A": "algebra",
    "T": "trigonometria",
    "G": "geometria",
    "C": "calculo",
}

CUESTIONARIO_CODES = {
    1: [f"{letra}{n}" for letra in "ATGC" for n in range(1, 11)],
    2: [f"{letra}{n}" for letra in "ATGC" for n in range(11, 21)],
}

C1_ACIERTOS_COL = {
    "algebra": "aciertos_c1_algebra",
    "trigonometria": "aciertos_c1_trigonometria",
    "geometria": "aciertos_c1_geometria",
    "calculo": "aciertos_c1_calculo",
}

C2_ACIERTOS_COL = {
    "algebra": "aciertos_c2_algebra",
    "trigonometria": "aciertos_c2_trigonometria",
    "geometria": "aciertos_c2_geometria",
    "calculo": "aciertos_c2_calculo",
}

USUARIO_RE = re.compile(r"^([A-Za-z]{2,5})[_\-\s]?\s*(\d+)$")


def _valor(row: tuple, idx: int | None) -> object | None:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _header_text(raw: object | None) -> str:
    if raw is None:
        return ""
    return str(raw).strip()


def detectar_columnas_cuestionario(headers: list) -> dict:
    """Detecta columnas por el texto del encabezado (posiciones variables)."""
    cols = {
        "email": None,
        "folio": None,
        "usuario": None,
        "comercial": None,
        "institucional": None,
        "num_cuenta": None,
        "score": None,
        "timestamp": None,
        "nombre": None,
        "preguntas": {},
    }

    for idx, raw in enumerate(headers):
        s = _header_text(raw)
        low = s.lower()
        if not low:
            continue

        matches = list(re.finditer(r"\(([ATGC])(\d+)\)", s))
        if matches:
            m = matches[-1]
            code = f"{m.group(1)}{int(m.group(2))}"
            if code not in cols["preguntas"]:
                cols["preguntas"][code] = idx
            continue

        if "comercial" in low:
            if cols["comercial"] is None:
                cols["comercial"] = idx
        elif "institucional" in low:
            if cols["institucional"] is None:
                cols["institucional"] = idx
        elif "email" in low or "correo" in low:
            if cols["email"] is None:
                cols["email"] = idx
        elif "cuenta" in low:
            if cols["num_cuenta"] is None:
                cols["num_cuenta"] = idx
        elif "folio" in low:
            if cols["folio"] is None:
                cols["folio"] = idx
        elif "usuario" in low:
            if cols["usuario"] is None:
                cols["usuario"] = idx
        elif "score" in low:
            if cols["score"] is None:
                cols["score"] = idx
        elif "timestamp" in low or "marca" in low or "fecha" in low or "hora" in low:
            if cols["timestamp"] is None:
                cols["timestamp"] = idx
        elif "nombre" in low:
            if cols["nombre"] is None:
                cols["nombre"] = idx

    if cols["email"] is None and len(headers) > 1:
        cols["email"] = 1
    if cols["folio"] is None and len(headers) > 3:
        cols["folio"] = 3
    if cols["usuario"] is None and len(headers) > 5:
        cols["usuario"] = 5
    if cols["timestamp"] is None and len(headers) > 0:
        cols["timestamp"] = 0

    return cols


def parse_usuario(usuario: object | None) -> tuple[str | None, int | None]:
    """Ej. 'ICO_118' -> ('ICO', 118). Clave de ingeniería + lugar de admisión."""
    if usuario is None:
        return None, None
    s = str(usuario).strip().upper()
    m = USUARIO_RE.match(s)
    if m:
        return m.group(1), int(m.group(2))
    return None, None


def _leer_workbook(file_bytes: bytes) -> tuple[list[str], list[tuple]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    if header_row is None:
        raise ValueError("El archivo no tiene fila de encabezados")

    headers = [_header_text(h) for h in header_row]
    return headers, rows


def _correos_fila(row: tuple, cols: dict) -> list[str]:
    emails: list[str] = []
    for key in ("email", "comercial", "institucional"):
        raw = _valor(row, cols[key])
        if raw is None:
            continue
        em = clean_email(raw)
        if em and em not in emails:
            emails.append(em)
    return emails


def _identificadores_fila(row: tuple, cols: dict) -> tuple[list[str], list[str]]:
    cuentas: list[str] = []
    folios: list[str] = []

    raw = _valor(row, cols["folio"])
    tipo, valor = clasificar_identificador(raw)
    if tipo == "cuenta":
        cuentas.append(valor)
    elif tipo == "folio":
        folios.append(valor)

    raw_cuenta = _valor(row, cols["num_cuenta"])
    if raw_cuenta is not None:
        cuenta = normalizar_cuenta(raw_cuenta)
        if cuenta:
            cuentas.append(cuenta)

    return cuentas, folios


def _obtener_timestamp(row: tuple, cols: dict) -> object | None:
    return _valor(row, cols["timestamp"])


def _respuestas_key_default() -> dict[str, str]:
    flat: dict[str, str] = {}
    for por_materia in DEFAULT_RESPUESTAS_CUESTIONARIO.values():
        flat.update(por_materia)
    return flat


def calcular_aciertos_cuestionario(
    row: tuple,
    preguntas: dict[str, int],
    respuestas_key: dict[str, str],
    codes: list[str],
) -> tuple[dict[str, int], list[dict]]:
    aciertos = {"algebra": 0, "trigonometria": 0, "geometria": 0, "calculo": 0}
    detalle: list[dict] = []
    for code in codes:
        idx = preguntas.get(code)
        if idx is None:
            detalle.append({"codigo": code, "respuesta": "", "correcta": False})
            continue
        raw = _valor(row, idx)
        respuesta = extract_answer_key_from_raw(raw)
        correcta = bool(respuesta) and respuesta == respuestas_key.get(code, "")
        if correcta:
            aciertos[MATERIA_FOR_CODE[code[0]]] += 1
        detalle.append(
            {
                "codigo": code,
                "respuesta": respuesta or "",
                "correcta": correcta,
            }
        )
    return aciertos, detalle


def _filtro_preguntas(cols: dict, cuestionario: int) -> dict[str, int]:
    allowed = set(CUESTIONARIO_CODES[cuestionario])
    return {code: idx for code, idx in cols["preguntas"].items() if code in allowed}


async def _upsert_cuestionario(
    db: AsyncSession,
    alumno_id: int,
    periodo: str,
    cuestionario: int,
    aciertos: dict[str, int],
    respuestas_json: str,
) -> None:
    result = await db.execute(
        select(ResultadoCuestionarioDiagnostico).where(
            ResultadoCuestionarioDiagnostico.alumno_id == alumno_id,
            ResultadoCuestionarioDiagnostico.periodo == periodo,
        )
    )
    row = result.scalars().first()
    if row is None:
        row = ResultadoCuestionarioDiagnostico(alumno_id=alumno_id, periodo=periodo)
        db.add(row)

    cols_map = C1_ACIERTOS_COL if cuestionario == 1 else C2_ACIERTOS_COL
    for materia, col in cols_map.items():
        setattr(row, col, aciertos[materia])
    if cuestionario == 1:
        row.respuestas_c1 = respuestas_json
    else:
        row.respuestas_c2 = respuestas_json

    await db.flush()


async def _indice_usuario(
    db: AsyncSession, alumno_details: dict[int, dict]
) -> dict[tuple[str, int], list[int]]:
    result = await db.execute(select(Alumno.id, Alumno.lugar_admision))
    lugar_map = dict(result.all())

    index: dict[tuple[str, int], list[int]] = {}
    for aid, info in alumno_details.items():
        clave = (info.get("ingenieria") or "").upper()
        lugar = lugar_map.get(aid)
        if clave and lugar is not None:
            index.setdefault((clave, lugar), []).append(aid)
    return index


def _detalle_resultado(
    bad: dict, alumno_id: int, aciertos: dict[str, int], detalle: list[dict]
) -> dict:
    info = bad[alumno_id]
    return {
        "alumno_id": alumno_id,
        "nombre_completo": info["nombre"],
        "numero_cuenta": info["cuenta"],
        "numero_folio": info["folio"],
        "correo": info["correo"],
        "ingenieria": info["ingenieria"],
        "aciertos": aciertos,
        "respuestas": detalle,
    }


async def procesar_cuestionario(
    db: AsyncSession,
    file_bytes: bytes,
    cuestionario: int,
    periodo: str,
) -> dict:
    if cuestionario not in (1, 2):
        raise ValueError(f"Cuestionario no válido: {cuestionario}")

    headers, rows = _leer_workbook(file_bytes)
    cols = detectar_columnas_cuestionario(headers)
    preguntas = _filtro_preguntas(cols, cuestionario)
    respuestas_key = _respuestas_key_default()

    email_map, cuenta_map, folio_map, alumno_details = await load_all_alumnos(db)
    usuario_index = await _indice_usuario(db, alumno_details)

    rango = await obtener_rango_periodo(db, periodo)

    encontrados_cola: list[dict] = []
    no_encontrados: list[dict] = []
    omitidas = 0

    for idx, row in enumerate(rows):
        ts = _obtener_timestamp(row, cols)
        if not fila_corresponde_periodo(ts, rango):
            omitidas += 1
            continue

        emails = _correos_fila(row, cols)
        cuentas, folios = _identificadores_fila(row, cols)
        usuario_raw = _valor(row, cols["usuario"])
        nombre = normalize_name(_valor(row, cols["nombre"]))

        alumno_id = None
        for em in emails:
            if em in email_map:
                alumno_id = email_map[em]
                break
        if alumno_id is None:
            for c in cuentas:
                if c in cuenta_map:
                    alumno_id = cuenta_map[c]
                    break
        if alumno_id is None:
            for f in folios:
                if f in folio_map:
                    alumno_id = folio_map[f]
                    break

        if alumno_id is None:
            motivo = "Sin coincidencia de correo/cuenta/folio"
            candidatos: list[dict] = []
            clave, lugar = parse_usuario(usuario_raw)
            if clave and lugar is not None:
                ids = usuario_index.get((clave, lugar))
                if ids and len(ids) == 1:
                    info = alumno_details[ids[0]]
                    candidatos.append(
                        {
                            "alumno_id": ids[0],
                            "nombre": info["nombre"],
                            "cuenta": info["cuenta"],
                            "correo": info["correo"],
                            "sugerido": True,
                            "motivo": "Coincidencia por usuario (carrera + lugar): confirmar",
                        }
                    )
                    motivo = (
                        "Sin coincidencia de correo/cuenta/folio; "
                        "usuario coincide con un solo alumno (carrera + lugar)"
                    )
            no_encontrados.append(
                {
                    "nombre_original": nombre,
                    "correo": emails[0] if emails else None,
                    "cuenta": cuentas[0] if cuentas else None,
                    "folio": folios[0] if folios else None,
                    "usuario": str(usuario_raw).strip()
                    if usuario_raw is not None
                    else None,
                    "cuestionario": cuestionario,
                    "motivo": motivo,
                    "candidatos": candidatos,
                    "indice": idx,
                }
            )
            continue

        encontrados_cola.append(
            {
                "idx": idx,
                "alumno_id": alumno_id,
                "row": row,
                "ts": ts,
            }
        )

    encontrados_cola.sort(key=lambda e: (ts_sort_key(e["ts"]), e["idx"]))

    vistos: set[int] = set()
    repetidos = 0
    resultados: list[dict] = []

    for e in encontrados_cola:
        if e["alumno_id"] in vistos:
            repetidos += 1
            continue
        vistos.add(e["alumno_id"])

        aciertos, detalle = calcular_aciertos_cuestionario(
            e["row"], preguntas, respuestas_key, CUESTIONARIO_CODES[cuestionario]
        )
        respuestas_json = json.dumps(detalle, ensure_ascii=False)
        await _upsert_cuestionario(
            db, e["alumno_id"], periodo, cuestionario, aciertos, respuestas_json
        )
        resultados.append(
            _detalle_resultado(alumno_details, e["alumno_id"], aciertos, detalle)
        )

    await db.commit()

    return {
        "cuestionario": cuestionario,
        "periodo": periodo,
        "rango_fechas": rango_fechas_dict(rango),
        "total_filas": len(rows),
        "encontrados": len(resultados),
        "no_encontrados": len(no_encontrados),
        "omitidas_otro_periodo": omitidas,
        "intentos_repetidos_ignorados": repetidos,
        "resultados": resultados,
        "no_encontrados_detalle": no_encontrados,
    }


async def corregir_matching_cuestionario(
    db: AsyncSession,
    file_bytes: bytes,
    cuestionario: int,
    periodo: str,
    correcciones: list[dict],
) -> dict:
    if cuestionario not in (1, 2):
        raise ValueError(f"Cuestionario no válido: {cuestionario}")

    headers, rows = _leer_workbook(file_bytes)
    cols = detectar_columnas_cuestionario(headers)
    preguntas = _filtro_preguntas(cols, cuestionario)
    respuestas_key = _respuestas_key_default()

    rango = await obtener_rango_periodo(db, periodo)

    _, cuenta_map, folio_map, alumno_details = await load_all_alumnos(db)

    correction_map = {c["indice"]: c["alumno_id"] for c in correcciones}
    resultados: list[dict] = []
    omitidas = 0

    for idx, row in enumerate(rows):
        ts = _obtener_timestamp(row, cols)
        if not fila_corresponde_periodo(ts, rango):
            omitidas += 1
            continue
        if idx not in correction_map:
            continue

        alumno_id = correction_map[idx]
        if alumno_id not in alumno_details:
            continue

        aciertos, detalle = calcular_aciertos_cuestionario(
            row, preguntas, respuestas_key, CUESTIONARIO_CODES[cuestionario]
        )
        respuestas_json = json.dumps(detalle, ensure_ascii=False)
        await _upsert_cuestionario(
            db, alumno_id, periodo, cuestionario, aciertos, respuestas_json
        )
        resultados.append(
            _detalle_resultado(alumno_details, alumno_id, aciertos, detalle)
        )

    await db.commit()

    return {
        "cuestionario": cuestionario,
        "periodo": periodo,
        "total_filas": len(resultados),
        "encontrados": len(resultados),
        "no_encontrados": 0,
        "omitidas_otro_periodo": omitidas,
        "intentos_repetidos_ignorados": 0,
        "resultados": resultados,
        "no_encontrados_detalle": [],
    }
