"""Carga del catálogo de Control Escolar (hoja "Datos catalogo" del CREANI).

Columnas (sin encabezados): A correo, B número de cuenta, C cuenta con "A"
adelante (se ignora), E nombre(s), F apellido paterno, G apellido materno,
I correo (repetido).

El catálogo no crea alumnos: solo completa a los que ya existen. Se buscan
SOLO por correo o por cuenta, nunca por nombre (hay homónimos).
- Provisional: toma nombre, apellidos y cuenta del catálogo (la cuenta solo si
  no es de otro alumno). Sigue siendo provisional hasta que llegue el padrón.
- Normal: solo se llenan la cuenta o el nombre si están vacíos.
- En ambos casos los correos nuevos se guardan como identificadores.
"""

from __future__ import annotations

import io
import unicodedata

import openpyxl
import xlrd
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.alumno import Alumno
from app.models.user import User
from app.services.alumnos_provisionales import nombre_vacio
from app.services.identificadores_service import RegistroIdentificadores
from app.services.normalizacion import normalizar_correo, normalizar_cuenta
from app.services.upload_diagnostico_service import load_all_alumnos

HOJA_CATALOGO = "datos catalogo"

COL_CORREO = 0
COL_CUENTA = 1
COL_NOMBRE = 4
COL_APELLIDO_PATERNO = 5
COL_APELLIDO_MATERNO = 6
COL_CORREO_2 = 8

FUENTE = "catalogo_control_escolar"


class CatalogoError(ValueError):
    """El archivo no tiene la hoja del catálogo."""


def _clave_hoja(nombre: str) -> str:
    s = unicodedata.normalize("NFKD", nombre)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _texto(valor: object | None) -> str:
    return " ".join(str(valor or "").split()).title()


def _celda(fila: tuple, col: int) -> object | None:
    return fila[col] if col < len(fila) else None


def leer_filas_catalogo(file_bytes: bytes, nombre_archivo: str) -> list[tuple]:
    """Filas crudas de la hoja "Datos catalogo" (.xlsx o .xls)."""
    if nombre_archivo.lower().endswith(".xls"):
        wb = xlrd.open_workbook(file_contents=file_bytes)
        hojas = {_clave_hoja(s.name): s for s in wb.sheets()}
        hoja = hojas.get(HOJA_CATALOGO)
        if hoja is None:
            raise CatalogoError(_mensaje_sin_hoja([s.name for s in wb.sheets()]))
        return [tuple(hoja.row_values(i)) for i in range(hoja.nrows)]

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        hojas = {_clave_hoja(n): n for n in wb.sheetnames}
        nombre = hojas.get(HOJA_CATALOGO)
        if nombre is None:
            raise CatalogoError(_mensaje_sin_hoja(wb.sheetnames))
        return list(wb[nombre].iter_rows(values_only=True))
    finally:
        wb.close()


def _mensaje_sin_hoja(hojas: list[str]) -> str:
    return (
        'El archivo no tiene la hoja "Datos catalogo". '
        f"Hojas encontradas: {', '.join(hojas) or 'ninguna'}"
    )


def parsear_fila(fila: tuple) -> dict | None:
    """Datos normalizados de una fila, o None si no trae correo ni cuenta."""
    correos = []
    for col in (COL_CORREO, COL_CORREO_2):
        correo = normalizar_correo(_celda(fila, col))
        if correo and correo not in correos:
            correos.append(correo)
    cuenta = normalizar_cuenta(_celda(fila, COL_CUENTA))
    if not correos and not cuenta:
        return None
    return {
        "correos": correos,
        "cuenta": cuenta,
        "nombre": _texto(_celda(fila, COL_NOMBRE)),
        "apellido_paterno": _texto(_celda(fila, COL_APELLIDO_PATERNO)),
        "apellido_materno": _texto(_celda(fila, COL_APELLIDO_MATERNO)),
    }


def emparejar(
    correos: list[str],
    cuenta: str | None,
    email_map: dict[str, int],
    cuenta_map: dict[str, int],
) -> tuple[int | None, dict[str, int]]:
    """Busca al alumno solo por correo y cuenta.

    Devuelve (alumno_id, coincidencias). Si los datos apuntan a dos alumnos
    distintos, alumno_id es None y `coincidencias` dice quién tiene cada dato.
    """
    coincidencias: dict[str, int] = {}
    for correo in correos:
        if correo in email_map:
            coincidencias[correo] = email_map[correo]
    if cuenta and cuenta in cuenta_map:
        coincidencias[cuenta] = cuenta_map[cuenta]
    ids = set(coincidencias.values())
    if len(ids) == 1:
        return ids.pop(), coincidencias
    return None, coincidencias


async def procesar_catalogo(
    db: AsyncSession, file_bytes: bytes, nombre_archivo: str
) -> dict:
    filas = leer_filas_catalogo(file_bytes, nombre_archivo)

    email_map, cuenta_map, folio_map, alumno_details = await load_all_alumnos(db)
    registro = RegistroIdentificadores(
        db, email_map, cuenta_map, folio_map, alumno_details, fuente=FUENTE
    )

    total_filas = 0
    actualizados: list[dict] = []
    sin_cambios = 0
    no_encontrados: list[dict] = []
    conflictos: list[dict] = []

    for i, raw in enumerate(filas):
        datos = parsear_fila(raw)
        if datos is None:
            continue
        total_filas += 1
        fila_num = i + 1
        nombre_fila = " ".join(
            p for p in (datos["nombre"], datos["apellido_paterno"], datos["apellido_materno"]) if p
        )

        alumno_id, coincidencias = emparejar(
            datos["correos"], datos["cuenta"], email_map, cuenta_map
        )
        if alumno_id is None:
            base = {
                "fila": fila_num,
                "nombre": nombre_fila,
                "correos": datos["correos"],
                "cuenta": datos["cuenta"],
            }
            if coincidencias:
                conflictos.append({
                    **base,
                    "motivo": "El correo y la cuenta son de alumnos distintos",
                    "alumnos": [
                        {
                            "dato": dato,
                            "alumno_id": aid,
                            "nombre": alumno_details.get(aid, {}).get("nombre", "").strip(),
                        }
                        for dato, aid in coincidencias.items()
                    ],
                })
            else:
                no_encontrados.append(base)
            continue

        detalle = alumno_details[alumno_id]
        provisional = bool(detalle.get("provisional"))
        cambios: list[str] = []

        alumno = await db.get(Alumno, alumno_id)
        user = await db.get(User, alumno.usuario_id)

        # Nombre: el provisional toma el del catálogo; el normal solo si no tiene.
        if datos["nombre"] and (provisional or nombre_vacio(detalle.get("nombre"))):
            nuevo = (datos["nombre"], datos["apellido_paterno"], datos["apellido_materno"])
            if nuevo != (user.nombre, user.apellido_paterno, user.apellido_materno):
                user.nombre, user.apellido_paterno, user.apellido_materno = nuevo
                detalle["nombre"] = f"{nuevo[0]} {nuevo[1]} {nuevo[2]}"
                cambios.append("nombre")

        # Cuenta del provisional aunque ya tuviera otra (si no es de otro alumno).
        # Al normal sin cuenta se la asigna RegistroIdentificadores.
        cuenta = datos["cuenta"]
        if (
            provisional
            and cuenta
            and alumno.numero_cuenta != cuenta
            and cuenta_map.get(cuenta) in (None, alumno_id)
        ):
            alumno.numero_cuenta = cuenta
            cuenta_map[cuenta] = alumno_id
            detalle["cuenta"] = cuenta
            cambios.append("cuenta")

        antes = registro.nuevos
        await registro.registrar(
            alumno_id,
            correos=datos["correos"],
            cuentas=[cuenta] if cuenta else [],
            indice=fila_num,
        )
        if registro.nuevos > antes:
            cambios.append("identificadores")

        if cambios:
            actualizados.append({
                "fila": fila_num,
                "alumno_id": alumno_id,
                "nombre": detalle["nombre"].strip(),
                "provisional": provisional,
                "cambios": cambios,
            })
        else:
            sin_cambios += 1

    return {
        "total_filas": total_filas,
        "actualizados": len(actualizados),
        "sin_cambios": sin_cambios,
        "no_encontrados": len(no_encontrados),
        "conflictos": conflictos,
        "actualizados_detalle": actualizados,
        "no_encontrados_detalle": no_encontrados,
        **registro.resumen(),
    }
