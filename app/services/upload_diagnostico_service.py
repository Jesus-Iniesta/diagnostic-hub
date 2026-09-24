from __future__ import annotations

import difflib
import io
import json
import re
import unicodedata
from datetime import date, datetime

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.alumno import Alumno
from app.models.user import User
from app.repositories.diagnostico_repository import DiagnosticoRepository
from app.seeds.data.respuestas_diagnostico import DEFAULT_RESPUESTAS
from app.services.normalizacion import (
    clasificar_identificador,
    normalizar_cuenta,
    normalizar_folio,
    normalizar_nombre,
)


MATERIA_COLUMNS: dict[str, dict[str, int]] = {
    "algebra": {"prefix": "FA", "start_col": 7, "questions": 20},
    "trigonometria": {"prefix": "FT", "start_col": 7, "questions": 20},
    "geometria": {"prefix": "FG", "start_col": 7, "questions": 20},
    "calculo": {"prefix": "FC", "start_col": 7, "questions": 20},
}

COL_TIMESTAMP = 0
COL_EMAIL = 1
COL_SCORE = 2
COL_NAME = 3
COL_CUENTA = 4
COL_FOLIO = 5
COL_HOUR = 6


def extraer_ano_periodo(periodo: str) -> int | None:
    if not periodo:
        return None
    m = re.match(r"\s*(\d{4})", str(periodo))
    if m:
        return int(m.group(1))
    return None


def extraer_ano_timestamp(raw: object | None) -> int | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.year
    if isinstance(raw, date):
        return raw.year
    s = str(raw).strip()
    if not s:
        return None
    m = re.search(r"\b(19|20)\d{2}\b", s)
    if m:
        return int(m.group(0))
    return None


def fila_corresponde_periodo(raw_timestamp: object | None, periodo: str) -> bool:
    ano_periodo = extraer_ano_periodo(periodo)
    ano_ts = extraer_ano_timestamp(raw_timestamp)
    if ano_periodo is None or ano_ts is None:
        return True
    return ano_ts == ano_periodo


def ts_sort_key(raw: object | None) -> tuple:
    """Clave para ordenar intentos por marca temporal (None va al final)."""
    if isinstance(raw, datetime):
        return (0, raw.replace(tzinfo=None))
    if isinstance(raw, date):
        return (0, datetime.combine(raw, datetime.min))
    if raw is None:
        return (2, datetime.min)
    s = str(raw).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
        "%m/%d/%Y %H:%M:%S",
        "%Y-%m-%d",
        "%d-%m-%Y %H:%M:%S",
    ):
        try:
            return (0, datetime.strptime(s, fmt))
        except ValueError:
            continue
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return (0, dt.replace(tzinfo=None))
    except ValueError:
        return (1, s)


def normalize_name(raw: str) -> str:
    if not raw:
        return ""
    s = str(raw).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def normalize_name_display(raw: str) -> str:
    if not raw:
        return ""
    s = str(raw).strip()
    s = re.sub(r"\s+", " ", s)
    nfkd = unicodedata.normalize("NFKD", s)
    without_accents = "".join(c for c in nfkd if not unicodedata.combining(c))
    return without_accents.title()


def normalize_name_for_match(raw: str) -> str:
    if not raw:
        return ""
    s = str(raw).strip().lower()
    s = re.sub(r"\s+", " ", s)
    nfkd = unicodedata.normalize("NFKD", s)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def normalize_email(raw: object | None) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().lower()
    return s if "@" in s else None


normalize_cuenta = normalizar_cuenta
normalize_folio = normalizar_folio


def clasificar_columnas_cuenta_folio(
    raw_cuenta: object | None,
    raw_folio: object | None,
) -> tuple[str | None, str | None]:
    tipo_cuenta, valor_cuenta = clasificar_identificador(raw_cuenta)
    tipo_folio, valor_folio = clasificar_identificador(raw_folio)

    cuenta: str | None = None
    folio: str | None = None

    if tipo_cuenta == "cuenta":
        cuenta = valor_cuenta
    elif tipo_cuenta == "folio" and tipo_folio != "folio":
        folio = valor_cuenta

    if tipo_folio == "folio":
        folio = valor_folio
    elif tipo_folio == "cuenta" and cuenta is None:
        cuenta = valor_folio

    return cuenta, folio


def extract_answer_key_from_raw(raw: object | None) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None

    s_lower = s.lower()

    m = re.search(r"la\s+opci[oó]n\s+([abcd])\)", s_lower)
    if m:
        return m.group(1)

    m = re.search(r"\(([abcd])\)", s_lower)
    if m:
        return m.group(1)

    m = re.search(r"\b([abcd])\)", s_lower)
    if m:
        return m.group(1)

    if s_lower in ("a", "b", "c", "d"):
        return s_lower

    num_map = {"1": "a", "2": "b", "3": "c", "4": "d"}
    if s in num_map:
        return num_map[s]

    if len(s) == 1 and s_lower in ("a", "b", "c", "d"):
        return s_lower

    nfkd = unicodedata.normalize("NFKD", s_lower)
    no_accents = "".join(c for c in nfkd if not unicodedata.combining(c))
    no_spaces = re.sub(r"\s+", "", no_accents)
    m = re.match(r"^opcion([1-4])$", no_spaces)
    if m:
        return num_map[m.group(1)]

    return None


def calculate_section_score(
    answers: list[str | None],
    correct_key: list[str],
    section_start: int,
    section_size: int = 5,
) -> float:
    correct = 0
    for i in range(section_start, min(section_start + section_size, len(answers))):
        if i < len(correct_key) and answers[i] == correct_key[i]:
            correct += 1
    return round(correct * 0.5, 2)


def calculate_exam_score(
    answers: list[str | None],
    correct_key: list[str],
) -> float:
    correct = 0
    for i in range(len(answers)):
        if i < len(correct_key) and answers[i] == correct_key[i]:
            correct += 1
    return round(correct * 0.5, 2)


def generate_answer_codes(prefix: str, count: int = 20) -> list[str]:
    codes = []
    for i in range(count):
        section = i // 5 + 1
        pos = i % 5 + 1
        codes.append(f"{prefix}{section}{pos}")
    return codes


async def load_all_alumnos(
    db: AsyncSession,
) -> tuple[
    dict[str, int],
    dict[str, int],
    dict[str, int],
    dict[str, dict],
]:
    email_map: dict[str, int] = {}
    cuenta_map: dict[str, int] = {}
    folio_map: dict[str, int] = {}
    alumno_details: dict[int, dict] = {}

    stmt = (
        select(Alumno, User)
        .join(User, Alumno.usuario_id == User.id)
        .options(selectinload(Alumno.ingenieria))
    )
    result = await db.execute(stmt)
    rows = result.all()

    for alumno, user in rows:
        if user.correo_personal:
            email_map[user.correo_personal.lower().strip()] = alumno.id
        if user.correo_institucional:
            email_map[user.correo_institucional.lower().strip()] = alumno.id
        if alumno.numero_cuenta:
            cuenta_map[alumno.numero_cuenta] = alumno.id
        if alumno.numero_folio:
            folio_map[alumno.numero_folio] = alumno.id
        alumno_details[alumno.id] = {
            "nombre": f"{user.nombre} {user.apellido_paterno} {user.apellido_materno}",
            "cuenta": alumno.numero_cuenta,
            "folio": alumno.numero_folio,
            "correo": user.correo_personal,
            "ingenieria": alumno.ingenieria.clave if alumno.ingenieria else "",
        }

    return email_map, cuenta_map, folio_map, alumno_details


def find_alumno(
    email: str | None,
    cuenta: str | None,
    folio: str | None,
    email_map: dict[str, int],
    cuenta_map: dict[str, int],
    folio_map: dict[str, int],
) -> int | None:
    if email and email in email_map:
        return email_map[email]
    if cuenta and cuenta in cuenta_map:
        return cuenta_map[cuenta]
    if folio and folio in folio_map:
        return folio_map[folio]
    return None


def find_candidates_legacy(
    name: str,
    alumno_details: dict[int, dict],
    limit: int = 5,
) -> list[dict]:
    if not name:
        return []
    normalized = normalize_name_for_match(name)
    candidates = []
    for aid, info in alumno_details.items():
        db_name = normalize_name_for_match(info["nombre"])
        if normalized in db_name or db_name in normalized:
            candidates.append({
                "alumno_id": aid,
                "nombre": info["nombre"],
                "cuenta": info["cuenta"],
                "correo": info["correo"],
            })
            if len(candidates) >= limit:
                break
    return candidates


def find_candidates(
    name: str,
    alumno_details: dict[int, dict],
    limit: int = 5,
) -> list[dict]:
    n = normalizar_nombre(name)
    if not n:
        return []

    exactos: list[dict] = []
    for aid, info in alumno_details.items():
        if normalizar_nombre(info["nombre"]) == n:
            exactos.append({
                "alumno_id": aid,
                "nombre": info["nombre"],
                "cuenta": info["cuenta"],
                "correo": info["correo"],
                "sugerido": False,
                "similitud": 100,
            })

    if exactos:
        if len(exactos) == 1:
            exactos[0]["sugerido"] = True
        return exactos

    candidatos = []
    for aid, info in alumno_details.items():
        db_n = normalizar_nombre(info["nombre"])
        if not db_n:
            continue
        ratio = difflib.SequenceMatcher(None, n, db_n).ratio()
        if ratio >= 0.80:
            candidatos.append({
                "alumno_id": aid,
                "nombre": info["nombre"],
                "cuenta": info["cuenta"],
                "correo": info["correo"],
                "sugerido": False,
                "similitud": int(round(ratio * 100)),
            })

    candidatos.sort(key=lambda c: c["similitud"], reverse=True)
    return candidatos[:limit]


async def procesar_examen_diagnostico(
    db: AsyncSession,
    file_bytes: bytes,
    materia: str,
    periodo: str,
) -> dict:
    if materia not in MATERIA_COLUMNS:
        raise ValueError(f"Materia no válida: {materia}")

    config = MATERIA_COLUMNS[materia]
    prefix = config["prefix"]
    question_count = config["questions"]

    repo = DiagnosticoRepository(db)
    answer_codes = generate_answer_codes(prefix, question_count)

    guardadas = await repo.get_respuestas_correctas(materia, periodo)
    if guardadas:
        respuestas_key = guardadas
        advertencia = None
    else:
        respuestas_key = dict(DEFAULT_RESPUESTAS[materia])
        advertencia = (
            f"No había respuestas correctas guardadas para el periodo {periodo}; "
            "se usó la clave por defecto."
        )

    correct_key = [respuestas_key.get(code, "").lower() for code in answer_codes]

    email_map, cuenta_map, folio_map, alumno_details = await load_all_alumnos(db)

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    resultados = []
    no_encontrados = []
    omitidas = 0
    encontrados = []

    for idx, row in enumerate(rows):
        raw_timestamp = row[COL_TIMESTAMP] if len(row) > COL_TIMESTAMP else None
        if not fila_corresponde_periodo(raw_timestamp, periodo):
            omitidas += 1
            continue

        raw_email = row[COL_EMAIL] if len(row) > COL_EMAIL else None
        raw_name = row[COL_NAME] if len(row) > COL_NAME else None
        raw_cuenta = row[COL_CUENTA] if len(row) > COL_CUENTA else None
        raw_folio = row[COL_FOLIO] if len(row) > COL_FOLIO else None

        email = normalize_email(raw_email)
        cuenta, folio = clasificar_columnas_cuenta_folio(raw_cuenta, raw_folio)
        nombre_original = normalize_name(str(raw_name) if raw_name else "")

        alumno_id = find_alumno(email, cuenta, folio, email_map, cuenta_map, folio_map)

        if alumno_id is None:
            candidates = find_candidates(nombre_original, alumno_details)
            if candidates and candidates[0]["sugerido"]:
                motivo = "Coincidencia exacta de nombre: confirmar"
            elif candidates and candidates[0]["similitud"] == 100:
                motivo = "Nombre ambiguo: varios alumnos con el mismo nombre"
            else:
                motivo = "Sin coincidencia de correo/cuenta/folio; revisar candidatos"
            no_encontrados.append({
                "nombre_original": nombre_original,
                "correo": email,
                "cuenta": cuenta,
                "folio": folio,
                "materia": materia,
                "motivo": motivo,
                "candidatos": candidates,
                "indice": idx,
            })
            continue

        encontrados.append({
            "idx": idx,
            "alumno_id": alumno_id,
            "row": row,
            "ts": raw_timestamp,
        })

    # El CREANI usa el PRIMER intento (marca temporal más antigua); los demás se ignoran.
    encontrados.sort(key=lambda e: (ts_sort_key(e["ts"]), e["idx"]))
    vistos: set[int] = set()
    repetidos = 0

    for e in encontrados:
        if e["alumno_id"] in vistos:
            repetidos += 1
            continue
        vistos.add(e["alumno_id"])

        row = e["row"]
        answers = []
        for qi in range(question_count):
            col_idx = config["start_col"] + qi
            raw_answer = row[col_idx] if len(row) > col_idx else None
            key = extract_answer_key_from_raw(raw_answer)
            answers.append(key)

        detail = alumno_details[e["alumno_id"]]
        respuestas_details = []
        correctas_count = 0
        for qi, code in enumerate(answer_codes):
            is_correct = answers[qi] is not None and answers[qi] == correct_key[qi]
            if is_correct:
                correctas_count += 1
            respuestas_details.append({
                "codigo": code,
                "respuesta": answers[qi] or "",
                "correcta": is_correct,
            })

        puntaje = calculate_exam_score(answers, correct_key)

        respuestas_json = json.dumps(respuestas_details, ensure_ascii=False)

        await repo.upsert_resultado(
            alumno_id=e["alumno_id"],
            periodo=periodo,
            respuestas_json=respuestas_json,
            materia=materia,
            puntaje=puntaje,
        )

        resultados.append({
            "alumno_id": e["alumno_id"],
            "nombre_completo": detail["nombre"],
            "numero_cuenta": detail["cuenta"],
            "numero_folio": detail["folio"],
            "correo": detail["correo"],
            "ingenieria": detail["ingenieria"],
            "respuestas": respuestas_details,
            "puntaje": puntaje,
        })

    await db.commit()

    return {
        "materia": materia,
        "periodo": periodo,
        "total_filas": len(rows),
        "encontrados": len(resultados),
        "no_encontrados": len(no_encontrados),
        "omitidas_otro_periodo": omitidas,
        "intentos_repetidos_ignorados": repetidos,
        "advertencia": advertencia,
        "resultados": resultados,
        "no_encontrados_detalle": no_encontrados,
    }


async def corregir_matching_diagnostico(
    db: AsyncSession,
    materia: str,
    periodo: str,
    file_bytes: bytes,
    correcciones: list[dict],
) -> dict:
    if materia not in MATERIA_COLUMNS:
        raise ValueError(f"Materia no válida: {materia}")

    config = MATERIA_COLUMNS[materia]
    prefix = config["prefix"]
    question_count = config["questions"]

    repo = DiagnosticoRepository(db)
    answer_codes = generate_answer_codes(prefix, question_count)

    guardadas = await repo.get_respuestas_correctas(materia, periodo)
    if guardadas:
        respuestas_key = guardadas
        advertencia = None
    else:
        respuestas_key = dict(DEFAULT_RESPUESTAS[materia])
        advertencia = (
            f"No había respuestas correctas guardadas para el periodo {periodo}; "
            "se usó la clave por defecto."
        )

    correct_key = [respuestas_key.get(code, "").lower() for code in answer_codes]

    _, cuenta_map, folio_map, alumno_details = await load_all_alumnos(db)

    correction_map = {c["indice"]: c["alumno_id"] for c in correcciones}

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    resultados = []
    omitidas = 0

    for idx, row in enumerate(rows):
        raw_timestamp = row[COL_TIMESTAMP] if len(row) > COL_TIMESTAMP else None
        if not fila_corresponde_periodo(raw_timestamp, periodo):
            omitidas += 1
            continue
        if idx not in correction_map:
            continue

        alumno_id = correction_map[idx]
        if alumno_id not in alumno_details:
            continue

        answers = []
        for qi in range(question_count):
            col_idx = config["start_col"] + qi
            raw_answer = row[col_idx] if len(row) > col_idx else None
            key = extract_answer_key_from_raw(raw_answer)
            answers.append(key)

        detail = alumno_details[alumno_id]
        respuestas_details = []
        for qi, code in enumerate(answer_codes):
            is_correct = answers[qi] is not None and answers[qi] == correct_key[qi]
            respuestas_details.append({
                "codigo": code,
                "respuesta": answers[qi] or "",
                "correcta": is_correct,
            })

        puntaje = calculate_exam_score(answers, correct_key)
        respuestas_json = json.dumps(respuestas_details, ensure_ascii=False)

        await repo.upsert_resultado(
            alumno_id=alumno_id,
            periodo=periodo,
            respuestas_json=respuestas_json,
            materia=materia,
            puntaje=puntaje,
        )

        resultados.append({
            "alumno_id": alumno_id,
            "nombre_completo": detail["nombre"],
            "numero_cuenta": detail["cuenta"],
            "numero_folio": detail["folio"],
            "correo": detail["correo"],
            "ingenieria": detail["ingenieria"],
            "respuestas": respuestas_details,
            "puntaje": puntaje,
        })

    await db.commit()

    return {
        "materia": materia,
        "periodo": periodo,
        "total_filas": len(resultados),
        "encontrados": len(resultados),
        "no_encontrados": 0,
        "omitidas_otro_periodo": omitidas,
        "advertencia": advertencia,
        "resultados": resultados,
        "no_encontrados_detalle": [],
    }


async def get_resultados_consolidados(
    db: AsyncSession, periodo: str
) -> list[dict]:
    repo = DiagnosticoRepository(db)
    resultados = await repo.get_resultados_by_periodo(periodo)

    alumno_ids = [r.alumno_id for r in resultados]
    if not alumno_ids:
        return []

    stmt = (
        select(Alumno, User)
        .join(User, Alumno.usuario_id == User.id)
        .where(Alumno.id.in_(alumno_ids))
        .options(selectinload(Alumno.ingenieria))
    )
    result = await db.execute(stmt)
    rows = result.all()

    alumno_map = {}
    for alumno, user in rows:
        alumno_map[alumno.id] = {
            "alumno_id": alumno.id,
            "nombre_completo": f"{user.nombre} {user.apellido_paterno} {user.apellido_materno}",
            "numero_cuenta": alumno.numero_cuenta,
            "numero_folio": alumno.numero_folio,
            "correo": user.correo_personal,
            "ingenieria": alumno.ingenieria.clave if alumno.ingenieria else "",
        }

    consolidated = []
    for r in resultados:
        info = alumno_map.get(r.alumno_id, {})
        consolidated.append({
            **info,
            "puntaje_algebra": r.puntaje_algebra,
            "puntaje_trigonometria": r.puntaje_trigonometria,
            "puntaje_geometria": r.puntaje_geometria,
            "puntaje_calculo": r.puntaje_calculo,
            "promedio_diagnostico": r.promedio_diagnostico,
        })

    return consolidated


def generate_excel_resultados(consolidated: list[dict], periodo: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados Diagnóstico"

    headers = [
        "Nombre", "No. Cuenta", "No. Folio", "Correo", "Ingeniería",
        "Álgebra", "Trigonometría", "Geometría", "Cálculo", "Promedio",
    ]
    ws.append(headers)

    for row in consolidated:
        ws.append([
            row.get("nombre_completo", ""),
            row.get("numero_cuenta", ""),
            row.get("numero_folio", ""),
            row.get("correo", ""),
            row.get("ingenieria", ""),
            row.get("puntaje_algebra"),
            row.get("puntaje_trigonometria"),
            row.get("puntaje_geometria"),
            row.get("puntaje_calculo"),
            row.get("promedio_diagnostico"),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
